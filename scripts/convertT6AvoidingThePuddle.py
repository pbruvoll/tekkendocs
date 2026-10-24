"""Converts the Tekken 6 frame data workbook (Tekken6FrameData.xlsm) to one csv per character.

The workbook has one sheet per character. Each sheet starts with the standard moves,
followed by one section per stance, where a section is a title row and a header row.
The sections are flattened into one table, and moves done from inside a stance get the
stance code as prefix, the same notation as the T7 frame data (e.g. "FLA 3").
The first rows of a stance section are the moves entering the stance, so they keep their
command as is.

The csv files use the same layout as the T7 sheets, so the site can parse them the same way.

The character avatars are pictures on the first sheet, each placed right above the cell with
the character name.

The workbook is the Tekken 6 BR data from Avoiding the Puddle, machine translated from a japanese wiki.
The frame data on the site is made by scripts/buildT6FrameData.py from better sources, which uses
these csv files to fill in frames missing from those sources.

Usage: python scripts/convertT6AvoidingThePuddle.py <path to Tekken6FrameData.xlsm>
Output: data/t6/atp/<character-id>.csv and data/t6/avatars/<character-id>.webp
"""

import csv
import io
import os
import posixpath
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from collections import Counter

import openpyxl
from PIL import Image, ImageOps

CSV_SEP = ";"
HEADERS = [
    "Command",
    "Hit level",
    "Damage",
    "Start up frame",
    "Block frame",
    "Hit frame",
    "Counter hit frame",
    "Notes",
]

repoRoot = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
outputRoot = os.path.join(repoRoot, "data", "t6", "atp")
avatarOutputRoot = os.path.join(repoRoot, "data", "t6", "avatars")
# the avatars in the workbook vary a few pixels in size, so they are all cropped to the same size
AVATAR_SIZE = (90, 117)
# some avatars have white lines along the edges, which are at most a few pixels wide
MAX_WHITE_EDGE = 6

xmlNamespaces = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

# sheet name in the workbook -> character id on the site. Mokujin has no data in the workbook
sheetToCharacterId = {
    "Alisa Bosconovitch": "alisa",
    "Anna Williams": "anna",
    "Armor King": "armor-king",
    "Asuka Kazama": "asuka",
    "Baek Doo San": "baek",
    "Bob Richards": "bob",
    "Bruce Irvin": "bruce",
    "Bryan Fury": "bryan",
    "Christie Monteiro": "christie",
    "Craig Marduk": "marduk",
    "Devil Jin": "devil-jin",
    "Eddy Gordo": "eddy",
    "Feng Wei": "feng",
    "Ganryu": "ganryu",
    "Heihachi Mishima": "heihachi",
    "Hwoarang": "hwoarang",
    "Jack-6": "jack-6",
    "Jin Kazama": "jin",
    "Julia Chang": "julia",
    "Kazuya Mishima": "kazuya",
    "King": "king",
    "Kuma": "kuma",
    "Lars Alexandersson": "lars",
    "Lee Chaolan": "lee",
    "Lei Wulong": "lei",
    "Leo Kliesen": "leo",
    "Lili Rochefort": "lili",
    "Ling Xiaoyu": "xiaoyu",
    "Marshall Law": "law",
    "Miguel Caballero Rojo": "miguel",
    "Nina Williams": "nina",
    "Paul Phoenix": "paul",
    "Raven": "raven",
    "Roger Jr.": "roger-jr",
    "Sergei Dragunov": "dragunov",
    "Steve Fox": "steve",
    "Wang Jinrei": "wang",
    "Yoshimitsu": "yoshimitsu",
    "Zafina": "zafina",
}

# stance section title -> (stance code, number of leading rows entering the stance).
# A code of None means the commands in the section are already complete (e.g. crouch dash
# moves written as f,n,d,df+1), so they are kept as they are.
# Sections not listed are treated as standard moves.
stanceSections = {
    "Crouch Dash": (None, 0),
    "Back Sway": (None, 0),
    "Mist Step": (None, 0),
    "Back Turned": ("BT", 0),
    "Single Boot": ("SBT", 2),
    "Destroy Form": ("DES", 1),
    "Dual Boot": ("DBT", 1),
    "Chaos Judgment": ("CHJ", 1),
    "Haze Palm": (None, 0),
    "Leg Cutter": (None, 0),
    "Flamingo": ("FLA", 1),
    "Spinning Ball": ("SPB", 1),
    "Muay Thai": ("MT", 1),
    "Taunt": ("TNT", 1),
    "Rewinder": (None, 0),
    "Handstand": ("HSP", 4),
    "Relax": ("RLX", 1),
    "Vale Tudo": ("VTS", 2),
    "Deity Flight": ("FLY", 1),
    "Shifting Clouds": (None, 0),
    "Kempo": (None, 0),
    "Sumo Sit": ("SUMO", 2),
    "Right Foot Forward": ("RFF", 0),
    "Left Flamingo": ("LFS", 2),
    "Right Flamingo": ("RFS", 1),
    "Sitdown": ("SIT", 2),
    "Crouching Demon": ("CDS", 1),
    "Evasive Spin": ("ESP", 2),
    "Hawk Lunge": (None, 0),
    "Jaguar Step": (None, 0),
    "Hunting Bear": ("HBS", 1),
    "Prone": ("PRN", 1),
    "Bear Sit": ("SIT", 2),
    "Silent Entry": ("SEN", 1),
    "Dynamic Entry": ("DEN", 1),
    "Avalanche": (None, 0),
    "Hitman": ("HMS", 1),
    "Hypnotist": ("HYP", 1),
    "Art Of Phoenix": ("AOP", 1),
    "Savage": ("SAV", 1),
    "Haze": ("HAZ", 1),
    "Rocket": (None, 0),
    "Animal Kick": (None, 0),
    "Left Sway": ("LSW", 1),
    "Right Sway": ("RSW", 1),
    "Albatross": ("ALB", 2),
    "Ducking": ("DCK", 1),
    "Peekaboo": ("PKB", 1),
    "Sway": ("SWY", 1),
    "Flicker": ("FLK", 2),
    "Power Consumption": ("PWC", 2),
    "Tarantula": ("TRT", 2),
    "Scare Crow": ("SCR", 2),
    "Mantis": ("MNT", 2),
    # the titles below are machine translated from a japanese wiki, so they are matched
    # on their anchor name, e.g. "... & aname (suiho)"
    "aname (bat)": (None, 0),  # Yoshimitsu standard moves, sword drawn
    "aname (kamae)": ("KAM", 1),
    "aname (step)": (None, 0),
    "aname (suiho)": ("DRU", 1),
    "aname (sleep)": ("PLD", 0),
    "aname (atmh)": ("FDFA", 1),
    "aname (ash)": ("FDFT", 1),
    "aname (atma)": ("FUFA", 1),
    "aname (asa)": ("FUFT", 1),
    "aname (back)": ("BT", 0),
    "aname (hisui)": ("JAD", 1),
    "aname (hebi)": ("SNK", 1),
    "aname (ryu)": ("DRG", 1),
    "aname (hyo)": ("PAN", 1),
    "aname (tora)": ("TGR", 1),
    "aname (tsuru)": ("CRA", 1),
    "aname (kinkei)": ("KNK", 2),
    "aname (bokuho)": ("BOK", 1),
    "aname (sinpo)": (None, 0),
    "aname (cd)": ("DCH", 1),
    "aname (fake)": ("FKS", 1),
    "aname (kin)": ("KIN", 1),
    "aname (muso)": ("MED", 1),
    "aname (jirai)": ("JIR", 1),
    "aname (agura)": ("IND", 1),
    "aname (tonbo)": ("DGF", 1),
    "aname (mut)": ("NSS", 2),
}

# english names for the stances with machine translated titles, shown as a key for the stance codes
translatedStanceNames = {
    "aname (kamae)": "Kamae",
    "aname (suiho)": "Drunken Walk",
    "aname (sleep)": "Play Dead",
    "aname (atmh)": "Play Dead, face down, feet away",
    "aname (ash)": "Play Dead, face down, feet towards",
    "aname (atma)": "Play Dead, face up, feet away",
    "aname (asa)": "Play Dead, face up, feet towards",
    "aname (back)": "Back Turned",
    "aname (hisui)": "Jade",
    "aname (hebi)": "Snake",
    "aname (ryu)": "Dragon",
    "aname (hyo)": "Panther",
    "aname (tora)": "Tiger",
    "aname (tsuru)": "Crane",
    "aname (kinkei)": "Kinkei",
    "aname (bokuho)": "Bokuho",
    "aname (cd)": "Dragon Charge",
    "aname (fake)": "Fake Step",
    "aname (kin)": "Kincho",
    "aname (muso)": "Meditation",
    "aname (jirai)": "Jirai",
    "aname (agura)": "Indian Stance",
    "aname (tonbo)": "Dragonfly",
    "aname (mut)": "No Sword Stance",
}

# (pattern, replacement) for machine translated phrases with a known notation, applied to all commands
translatedPhrases = [
    # the translation added a comma inside "run"
    (r"\bru,n\b", "run"),
    (r"\b(in the|inside|being) run\b", "WR"),
    # sidestep
    (r"\bright side while on the move\b", "SSR"),
    (r"\bleft side while on the move\b", "SSL"),
    (r"\b(while )?on the move next\b", "SS"),
    (r"\b(side while on the move|move in the next|next moving)\b", "SS"),
    (r"\b(lateral movement during the|while lateral movement|During the lateral movement)\b", "SS"),
    # hit and block conditions
    (r"\((when hit|only hits)\)", "(on hit)"),
    (r"\bwhen it hits the front\b", "on hit from the front"),
    (r"\bwhen it hits\b", "on hit"),
    (r"\+?(and )?when the guard hits\b", "on block"),
]

# moves which only hit a grounded opponent. The condition is moved to the notes column
groundedOpponentPattern = r"^(grounded opponent|opponent is down|Down in person|OTG|opponent down in just)\s+"
GROUNDED_OPPONENT_NOTE = "Hits grounded opponent"

# stance code -> (pattern, replacement) for the stance name repeated in the command after the stance code,
# e.g. "DRU in the random walk 1" => "DRU 1".
# The phrases "... in the preparation f or" seem to be a translation of "while in stance", not a forward input
stancePhrases = {
    "*": [(r"\bin \?\?|\?\? in\b", ""), (r"^the\b", "")],
    "DRU": [(r"\??\s*\b(in|during) (the|a) random walk\b", ""), (r"\? random walk\b", "")],
    "SNK": [(r"\bsnake in the preparation( f or)?", ""), (r"\bHolding the snake in the\b", "")],
    "DRG": [(r"\bin the stance of the dragon\b", ""), (r"\bdragon in the preparation( f or)?", "")],
    "TGR": [
        (r"\bTiger in the preparation( f or)?", ""),
        (r"\bin the stance of (the|a) tiger\b", ""),
        (r"\bin Tiger stance\b", ""),
    ],
    "CRA": [(r"\bin the stance of the crane\b", ""), (r"\bcrane stance in\b", "")],
    "PAN": [
        (r"\bLeopards in (the preparation( f or)?|stance)", ""),
        (r"(~ ~\+)?in a leopard (ready|Holding)\b", ""),
        (r"\bin the stance of a leopard\b", ""),
    ],
    "JAD": [(r"\bBirds Sleeping Princess in\b", ""), (r"\bSleeping Princess in birds\b", "")],
    "FDFA": [
        (r"~ head in the prone sleeping posture\b", ""),
        (r"\bproneness of the head ~ during stance sleep\b", ""),
    ],
    "FDFT": [
        (r"~ (leg|feet) in the prone sleeping posture\b", ""),
        (r"\?\? lying face down leg side\+?", ""),
        (r"\bready to face down the leg during sleep ~ ~", ""),
    ],
    "FUFA": [
        (r"\bsupine posture of the head while sleeping ~", ""),
        (r"\bin the stance of the head lying on his back\b", ""),
        (r"~ head in the supine sleeping posture\b", ""),
        (r"\bin the supine posture of the head lying\b", ""),
        (r"\bin the head supine sleeping posture\b", ""),
    ],
    "FUFT": [
        (r"~ leg side on his back in bed ready\b", ""),
        (r"\bin his back leg side sleeping posture\b", ""),
        (r"\bbacks the stance leg during sleep ~ ~", ""),
        (r"\bthe supine posture during sleep leg ~", ""),
    ],
    "PLD": [(r"\bwhile sleeping posture\b", "")],
    "BT": [
        (r"\b(while )?f or (the )?back (in|of)( the)?\b", ""),
        (r"\bwhile back f or( the)?\b", ""),
        (r"\btowards the back in\b", ""),
        (r"\bFC back way\b", "FC"),
    ],
    "DCH": [(r"\bCD (inside|during( the)?|into|in)\b", "")],
    "FKS": [(r"\bFS success (.*)", r"\1 (on success)"), (r"\bFS in( the)?\b", "")],
    "FLY": [(r"\b(flying in the air|air flight)\b", "")],
    "KNK": [(r"\b(in Kinkei|Kinkei in)\b", "")],
    "JIR": [
        (r"\??\s*\bblade (in )?mine( in)?\b", ""),
        (r"\bwhile cutting mine\b", ""),
        (r"\bmines in the blades?\b", ""),
    ],
    "IND": [(r"\b(to )?Akira Naka (to )?one thousand\b", ""), (r"\bswastika in AGRA\b", "")],
    "DGF": [(r"\b(dragonfly swastika in|swastika in (the )?dragonfly)\b", "")],
    "KAM": [(r"\b(while poised|stance (in|during))\b", "")],
    "RFF": [(r"\bin the right stance\b", "")],
    "VTS": [(r"\bin Kurauchingu\b", "FC")],
}

# the codes which can start a command, the stance codes and sidestep
prefixCodes = {code for code, _ in stanceSections.values() if code} | {"SS"}

# words which are part of the notation, so a command containing other words is likely garbled
notationWords = {"qcf", "qcb", "hold", "or"} | {code.lower() for code in prefixCodes}


def fail(message):
    print("ERROR: " + message)
    sys.exit(1)


def cellToString(value):
    return "" if value is None else str(value).strip()


def findStance(title):
    """Returns the stance code, the number of entry rows and the stance name for a section title."""
    anchor = re.search(r"aname \(\w+\)", title)
    key = anchor.group(0) if anchor else title
    if key not in stanceSections:
        return None
    code, entryRowCount = stanceSections[key]
    return code, entryRowCount, translatedStanceNames.get(key, title)


def cleanCommand(command):
    # the translation added commas after n inside words, e.g. "ran,dom" and "an,d"
    command = re.sub(r"(?<=[A-Za-z])n,(?=[A-Za-z])", "n", command)
    # "_" is used to separate alternative inputs, e.g. "b_f"
    command = re.sub(r"\s*_\s*", " or ", command)
    return re.sub(r"\s+", " ", command).strip()


def normalizeSpaces(command):
    return re.sub(r"\s+", " ", command).strip()


def translatePhrases(command):
    """Replaces machine translated phrases with notation. Returns the command and a note."""
    for pattern, replacement in translatedPhrases:
        command = re.sub(pattern, replacement, command)
    note = ""
    if re.search(groundedOpponentPattern, command):
        command = re.sub(groundedOpponentPattern, "", command)
        note = GROUNDED_OPPONENT_NOTE
    return normalizeSpaces(command), note


def removeStancePhrases(command):
    """Removes the stance name repeated after the stance code, e.g. "DRU in the random walk 1" => "DRU 1"."""
    code, _, rest = command.partition(" ")
    if not rest or code not in prefixCodes:
        return command
    for pattern, replacement in stancePhrases.get(code, []) + stancePhrases["*"]:
        rest = re.sub(pattern, replacement, rest, flags=re.IGNORECASE)
    # the removed phrases can leave a "~" or "," in front of the input
    rest = normalizeSpaces(rest).lstrip("~, ")
    return f"{code} {rest}" if rest else code


def cleanValue(value):
    return re.sub(r"\s*,\s*", ",", value)


def isGarbled(command):
    words = re.findall(r"[A-Za-z]{3,}", command)
    return any(word.lower() not in notationWords for word in words)


def isSectionTitle(row, nextRow):
    filledCells = [cell for cell in row if cell]
    return len(filledCells) == 1 and row[0] and nextRow and nextRow[0] == "Command"


def convertSheet(sheet):
    rows = [[cellToString(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    rows = [row for row in rows if any(row)]
    moves = []
    # stance code -> stance name, for the codes used as prefix
    stances = {}
    stance = (None, 0, None)
    rowInSection = 0
    for i, row in enumerate(rows):
        nextRow = rows[i + 1] if i + 1 < len(rows) else None
        if isSectionTitle(row, nextRow):
            sectionStance = findStance(row[0])
            if sectionStance is None:
                print(f"  WARNING: unknown section '{row[0]}' in {sheet.title}, treated as standard moves")
                sectionStance = (None, 0, None)
            stance = sectionStance
            if stance[0]:
                stances.setdefault(stance[0], stance[2])
            rowInSection = 0
            continue
        if row[0] == "Command" or not row[0] or i == 0:
            continue

        command, note = translatePhrases(cleanCommand(row[0]))
        code, entryRowCount, _ = stance
        if code and rowInSection >= entryRowCount and not command.startswith(code + " "):
            command = f"{code} {command}"
        command = removeStancePhrases(command)
        rowInSection += 1

        damage, startup, block, hit, counterHit = [cleanValue(v) for v in (row[1:6] + [""] * 5)[:5]]
        moves.append([command, "", damage, startup, block, hit, counterHit, note])
    return moves, stances


def writeCsv(characterId, moves, stances):
    os.makedirs(outputRoot, exist_ok=True)
    path = os.path.join(outputRoot, characterId + ".csv")
    with open(path, "w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file, delimiter=CSV_SEP, lineterminator="\n")
        writer.writerow(["#frames_normal"])
        writer.writerow(HEADERS)
        writer.writerows(moves)
        if stances:
            writer.writerow(["#stances"])
            writer.writerow(["Code", "Stance"])
            writer.writerows(stances.items())


def isWhiteLine(pixels):
    return sum(sum(pixel) for pixel in pixels) / (3 * len(pixels)) > 240


def trimWhiteEdges(image):
    """Removes the white lines along the edges of an image.
    A white line can have a dark line outside it, so everything up to the last white line is removed."""
    width, height = image.size
    pixels = image.load()
    edges = {
        "top": lambda i: [pixels[x, i] for x in range(width)],
        "bottom": lambda i: [pixels[x, height - 1 - i] for x in range(width)],
        "left": lambda i: [pixels[i, y] for y in range(height)],
        "right": lambda i: [pixels[width - 1 - i, y] for y in range(height)],
    }
    trim = {
        edge: max((i + 1 for i in range(MAX_WHITE_EDGE) if isWhiteLine(getLine(i))), default=0)
        for edge, getLine in edges.items()
    }
    return image.crop((trim["left"], trim["top"], width - trim["right"], height - trim["bottom"]))


def writeAvatars(workbookPath, characterSelectSheet):
    """Saves the pictures on the character select sheet, named by the character below them."""
    names = {
        (rowIndex, columnIndex): value
        for rowIndex, row in enumerate(characterSelectSheet.iter_rows(values_only=True))
        for columnIndex, value in enumerate(row)
        if value
    }
    drawingFolder = "xl/drawings"
    os.makedirs(avatarOutputRoot, exist_ok=True)
    with zipfile.ZipFile(workbookPath) as archive:
        relations = ET.fromstring(archive.read(f"{drawingFolder}/_rels/drawing1.xml.rels"))
        imagePaths = {
            relation.get("Id"): posixpath.normpath(posixpath.join(drawingFolder, relation.get("Target")))
            for relation in relations
        }
        drawing = ET.fromstring(archive.read(f"{drawingFolder}/drawing1.xml"))
        for anchor in drawing:
            row = int(anchor.find("xdr:from/xdr:row", xmlNamespaces).text)
            column = int(anchor.find("xdr:from/xdr:col", xmlNamespaces).text)
            characterId = sheetToCharacterId.get(names.get((row + 1, column)))
            if not characterId:
                continue
            imageId = anchor.find(".//a:blip", xmlNamespaces).get(f"{{{xmlNamespaces['r']}}}embed")
            with Image.open(io.BytesIO(archive.read(imagePaths[imageId]))) as image:
                avatar = ImageOps.fit(trimWhiteEdges(image.convert("RGB")), AVATAR_SIZE)
                avatar.save(os.path.join(avatarOutputRoot, characterId + ".webp"), quality=90)


def main():
    if len(sys.argv) != 2:
        fail("Usage: python scripts/convertT6AvoidingThePuddle.py <path to Tekken6FrameData.xlsm>")
    workbookPath = sys.argv[1]
    if not os.path.isfile(workbookPath):
        fail("Could not find " + workbookPath)

    workbook = openpyxl.load_workbook(workbookPath, read_only=True, data_only=True)
    writeAvatars(workbookPath, workbook.worksheets[0])
    for sheet in workbook.worksheets:
        characterId = sheetToCharacterId.get(sheet.title)
        if not characterId:
            print(f"Skipping sheet '{sheet.title}'")
            continue
        moves, stances = convertSheet(sheet)
        writeCsv(characterId, moves, stances)

        garbled = [move[0] for move in moves if isGarbled(move[0])]
        duplicates = [command for command, count in Counter(move[0] for move in moves).items() if count > 1]
        print(f"{characterId}: {len(moves)} moves, {len(garbled)} garbled, {len(duplicates)} duplicated commands")
        for command in garbled:
            print(f"  garbled: {command}")
        for command in duplicates:
            print(f"  duplicate: {command}")


if __name__ == "__main__":
    main()
